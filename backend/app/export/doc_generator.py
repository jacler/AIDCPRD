"""Generate 《项目建设技术方案书》 Word document — deterministic sections."""

from __future__ import annotations

import io
import tempfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

DISCLAIMER_FOOTER = "本方案基于公开数据生成，仅供技术参考"

_DATA_SOURCES = [
    "data/public_sources/gpus/*.json",
    "data/public_sources/switches/*.json",
    "data/public_sources/optics/transceiver_types.json",
    "data/public_sources/racks/standard_42u.json",
    "data/public_sources/servers/gpu_node_8x.json",
    "data/public_sources/compatibility/*.json",
]


@dataclass
class TechnicalProposalInput:
    project_name: str
    target_gpus: int
    scenario: str
    description: str | None = None
    topology: dict[str, Any] | None = None
    cost_breakdown: dict[str, float] | None = None
    bom_items: list[dict[str, Any]] = field(default_factory=list)
    multi_plans: list[dict[str, Any]] = field(default_factory=list)
    compliance_result: dict[str, Any] | None = None
    consultation_summary: list[str] = field(default_factory=list)
    topology_compute_svg: str | None = None
    topology_storage_svg: str | None = None
    cabinet_svg: str | None = None
    convergence_ratio: str = "1:1"


def _set_run_font(run, size_pt: int = 11, bold: bool = False) -> None:
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        _set_run_font(run, size_pt=16 if level == 1 else 14, bold=True)


def _add_paragraph(doc: Document, text: str, *, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run_font(run, bold=bold)


def _svg_to_png_bytes(svg: str) -> bytes | None:
    try:
        from svglib.svglib import svg2rlg
        from reportlab.graphics import renderPM

        drawing = svg2rlg(io.BytesIO(svg.encode("utf-8")))
        if drawing is None:
            return None
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
            path = tmp.name
        renderPM.drawToFile(drawing, path, fmt="PNG")
        return Path(path).read_bytes()
    except Exception:
        return None


def _add_svg_figure(doc: Document, svg: str | None, caption: str) -> None:
    if not svg:
        _add_paragraph(doc, f"[附图缺失] {caption}")
        return
    png = _svg_to_png_bytes(svg)
    if png:
        stream = io.BytesIO(png)
        doc.add_picture(stream, width=Inches(6.0))
        cap = doc.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cap.runs:
            _set_run_font(run, size_pt=9)
    else:
        _add_paragraph(doc, f"{caption}（SVG 渲染暂不可用，请通过设计器导出 SVG 附图）")


def _hidden_cost_note(item: dict[str, Any]) -> str:
    dim = str(item.get("cost_dimension") or "")
    mapping = {
        "COMPUTE": "算力隐性成本：机房空间/配电",
        "NETWORK": "网络隐性成本：光模块损耗与备件",
        "STORAGE": "存储隐性成本：数据迁移与备份",
        "SOFTWARE": "软件隐性成本：许可续费",
        "INFRA": "基础设施隐性成本：集成调试",
    }
    return mapping.get(dim, "运维与备件")


def _add_footer_disclaimer(doc: Document) -> None:
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(DISCLAIMER_FOOTER)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 100, 100)
    _set_run_font(run, size_pt=8)


def generate_technical_proposal_docx(data: TechnicalProposalInput) -> bytes:
    """Build Word document bytes for 《项目建设技术方案书》."""
    doc = Document()
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d")

    title = doc.add_heading(f"{data.project_name} — 项目建设技术方案书", level=0)
    for run in title.runs:
        _set_run_font(run, size_pt=18, bold=True)
    _add_paragraph(doc, f"生成日期：{now}（UTC）")
    _add_paragraph(doc, f"目标 GPU：{data.target_gpus} · 场景：{data.scenario}")

    _add_heading(doc, "一、项目背景", level=1)
    bg = data.description or (
        f"本项目面向 {data.scenario} 场景，规划建设 {data.target_gpus} 卡 GPU 集群，"
        "采用 Spine-Leaf 网络架构与独立存储平面。"
    )
    _add_paragraph(doc, bg)

    _add_heading(doc, "二、需求诊断记录", level=1)
    if data.consultation_summary:
        for i, line in enumerate(data.consultation_summary, 1):
            _add_paragraph(doc, f"{i}. {line}")
    else:
        _add_paragraph(doc, "（暂无结构化诊断记录，建议通过顾问式对话补充后重新导出）")

    _add_heading(doc, "三、拓扑设计", level=1)
    topo = data.topology or {}
    compute = topo.get("compute") or {}
    network = topo.get("network") or {}
    storage = topo.get("storage") or {}
    _add_paragraph(
        doc,
        f"计算节点 {compute.get('servers', '—')} 台 · GPU {compute.get('gpus', '—')} 张 · "
        f"Leaf {network.get('leaf_switches', '—')} / Spine {network.get('spine_switches', '—')} · "
        f"存储节点 {storage.get('nodes', '—')} · 收敛比 {data.convergence_ratio}",
    )
    _add_paragraph(doc, "3.1 计算网拓扑", bold=True)
    _add_svg_figure(doc, data.topology_compute_svg, "图 3-1 计算网（Parameter Plane）")
    _add_paragraph(doc, "3.2 存储网拓扑", bold=True)
    _add_svg_figure(doc, data.topology_storage_svg, "图 3-2 存储网（Storage Plane）")
    _add_paragraph(doc, "3.3 机柜排布", bold=True)
    _add_svg_figure(doc, data.cabinet_svg, "图 3-3 机柜 U 位与 PDU 排布")

    _add_heading(doc, "四、方案权衡分析", level=1)
    if data.multi_plans:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        for i, label in enumerate(["方案", "定位", "收敛比", "5年 TCO (万元)"]):
            hdr[i].text = label
        for plan in data.multi_plans:
            row = table.add_row().cells
            row[0].text = str(plan.get("plan_id", ""))
            row[1].text = str(plan.get("headline", ""))[:80]
            row[2].text = str(plan.get("convergence_ratio", ""))
            tco = plan.get("tco_5y") or 0
            row[3].text = f"{float(tco) / 10_000:.2f}"
            for cell in row:
                for p in cell.paragraphs:
                    for run in p.runs:
                        _set_run_font(run, size_pt=9)
        for plan in data.multi_plans:
            trade_offs = plan.get("trade_offs") or []
            if trade_offs:
                _add_paragraph(doc, f"【{plan.get('plan_id')}】权衡：", bold=True)
                for t in trade_offs:
                    _add_paragraph(doc, f"  ⚠️ {t}")
    else:
        _add_paragraph(doc, "（请先生成成本 breakdown 与多方案对比后重新导出）")

    _add_heading(doc, "五、BOM 清单", level=1)
    if data.bom_items:
        table = doc.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        for i, label in enumerate(["类别", "型号", "数量", "单价", "小计", "隐性成本"]):
            hdr[i].text = label
        for item in data.bom_items:
            row = table.add_row().cells
            row[0].text = str(item.get("category", ""))
            row[1].text = str(item.get("model", ""))
            row[2].text = str(item.get("quantity", ""))
            row[3].text = str(item.get("unit_price", ""))
            row[4].text = str(item.get("total_price", ""))
            row[5].text = _hidden_cost_note(item)
    else:
        _add_paragraph(doc, "（BOM 为空）")

    if data.cost_breakdown:
        total = data.cost_breakdown.get("total", 0)
        _add_paragraph(doc, f"五维成本合计：{total} 元（含算力/网络/存储/软件/基础设施）")

    _add_heading(doc, "六、合规自查表", level=1)
    compliance = data.compliance_result or {}
    checklist = compliance.get("security_checklist") or []
    if checklist:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        for i, label in enumerate(["检查项", "状态", "补救措施", "法规依据"]):
            hdr[i].text = label
        for item in checklist:
            row = table.add_row().cells
            row[0].text = str(item.get("title", ""))
            row[1].text = str(item.get("status_icon", item.get("status", "")))
            row[2].text = str(item.get("remediation", ""))[:120]
            row[3].text = str(item.get("regulation_ref", ""))
    else:
        _add_paragraph(doc, "（尚未执行合规评估，请在方案 PK 页运行合规检查）")

    appendix = compliance.get("appendix_markdown")
    if appendix:
        _add_paragraph(doc, "合规判定依据附录（摘要）：", bold=True)
        for line in str(appendix).splitlines()[:20]:
            if line.strip():
                _add_paragraph(doc, line.strip()[:200])

    _add_heading(doc, "七、数据来源声明", level=1)
    for src in _DATA_SOURCES:
        _add_paragraph(doc, f"· {src}")
    _add_paragraph(doc, "硬件参数均来自公开 JSON 数据源；价格字段为 null 时表示需用户补充渠道价。")

    _add_footer_disclaimer(doc)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
